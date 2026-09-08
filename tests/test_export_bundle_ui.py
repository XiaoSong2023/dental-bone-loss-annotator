"""Exercise the UI export job and old-project entry with synthetic fixtures."""

import csv
import json
import os
import zipfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np
import pytest
from PySide6.QtCore import QSettings
from PySide6.QtGui import QCloseEvent

from app.core.models import Pt
from app.core.project import Project, STATUS_DONE
from app.imaging.loader import imwrite_unicode
from app.ui import main_window


@pytest.fixture
def export_window(qtbot, tmp_path, monkeypatch):
    settings = QSettings(str(tmp_path / "prefs.ini"), QSettings.Format.IniFormat)
    monkeypatch.setattr(main_window, "QSettings", lambda *_: settings)
    messages, errors = [], []
    monkeypatch.setattr(main_window.QMessageBox, "information", lambda *a: messages.append(a[2]))
    monkeypatch.setattr(main_window.QMessageBox, "critical", lambda *a: errors.append(a[2]))
    monkeypatch.setattr(main_window.QMessageBox, "warning", lambda *a: errors.append(a[2]))
    project = Project.create(tmp_path / "病例项目")
    for i in range(2):
        path = tmp_path / f"00{i + 1}_工程夹具.png"
        imwrite_unicode(path, (np.arange(12000).reshape(100, 120) + i).astype(np.uint8))
        entry = project.import_paths([path])[0][0]
        ann = project.load_annotation(entry)
        ann.patient.age, ann.patient.sex = 61, "女"
        tooth = ann.add_tooth(16, (10, 10, 60, 80))
        tooth.site("mesial").cej = Pt(20.5, 25.25)
        project.save_annotation(entry, ann)
        project.set_status(entry, STATUS_DONE)
    window = main_window.MainWindow()
    qtbot.addWidget(window)
    window._load_project(project)
    window.open_case(project.cases[0])
    return window, messages, errors


@pytest.mark.parametrize("kind", ["csv", "excel", "json"])
def test_ui_exports_data_and_images_without_changing_completed_cases(
    export_window, qtbot, monkeypatch, kind
):
    window, messages, errors = export_window
    project = window.project
    before_images = [project.image_path(e).read_bytes() for e in project.cases]
    # A last unsaved manual adjustment must appear in both data and images.
    window.ann.teeth[0].site("mesial").cej = Pt(24.75, 31.5)
    window._mark_dirty()
    json_path = project.exports_dir / "当前病例.json"
    monkeypatch.setattr(main_window.QFileDialog, "getSaveFileName", lambda *_: (str(json_path), ""))
    getattr(window, "export_" + kind)()
    assert window._export_busy and not window.act_import.isEnabled()
    close = QCloseEvent()
    window.closeEvent(close)
    assert not close.isAccepted()
    qtbot.waitUntil(lambda: not window._export_busy, timeout=15000)
    assert not errors and "标注图片" in messages[-1]
    assert window.act_import.isEnabled() and window._autosave.isActive()
    archives = list(project.exports_dir.glob("*.zip"))
    assert len(archives) == 1
    with zipfile.ZipFile(archives[0]) as archive:
        assert archive.testzip() is None
        assert len([p for p in archive.namelist() if p.endswith(".png")]) == (1 if kind == "json" else 2)
        assert "index.html" in archive.namelist()
    if kind == "csv":
        assert {p.name for p in project.exports_dir.iterdir()} == {"summary.csv", "sites.csv", archives[0].name}
        with (project.exports_dir / "sites.csv").open(encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream))
        assert (rows[0]["CEJ_x"], rows[0]["CEJ_y"]) == ("24.75", "31.5")
        assert not any("框" in key for key in rows[0])
    elif kind == "excel":
        from openpyxl import load_workbook
        book = load_workbook(project.exports_dir / "汇总.xlsx")
        assert book.sheetnames == ["患者级汇总", "位点级明细"]
        book.close()
    else:
        assert json.loads(json_path.read_text("utf-8"))["export_kind"] == "landmark_coordinates"
    reopened = Project.open(project.root)
    assert all(e.status == STATUS_DONE for e in reopened.cases)
    assert [reopened.image_path(e).read_bytes() for e in reopened.cases] == before_images
    assert reopened.load_annotation(reopened.cases[0]).teeth[0].site("mesial").cej.as_tuple() == (24.75, 31.5)


def test_failed_image_export_restores_ui_and_preserves_previous_zip(export_window, qtbot):
    window, messages, errors = export_window
    out = window.project.exports_dir / "annotated_images.zip"
    out.write_bytes(b"previous export")
    window.project.cases[1].image = "missing.png"
    window.project.save()
    window.export_csv()
    qtbot.waitUntil(lambda: not window._export_busy, timeout=15000)
    assert errors and not messages
    assert out.read_bytes() == b"previous export"
    assert not (window.project.exports_dir / "summary.csv").exists()
    assert window.act_import.isEnabled() and window._autosave.isActive()


def test_explicit_old_project_open_keeps_done_group(export_window, monkeypatch):
    window, _, errors = export_window
    root = window.project.root
    window.save_current(silent=True)
    monkeypatch.setattr(main_window, "discover_existing_projects", lambda _: [root])
    window.open_existing_project()
    assert not errors and all(e.status == STATUS_DONE for e in window.project.cases)
    assert window.cases.lists[1].count() == 2


def test_project_dialog_cannot_overwrite_existing_data(export_window, monkeypatch, tmp_path):
    window, _, errors = export_window
    current = window.project
    before = current.project_file.read_bytes()
    monkeypatch.setattr(main_window.QFileDialog, "getExistingDirectory", lambda *_: str(current.root))
    window.new_project()
    assert errors and current.project_file.read_bytes() == before
    invalid = tmp_path / "坏项目"
    invalid.mkdir()
    (invalid / "project.json").write_text("broken", encoding="utf-8")
    window._open_project_path(invalid)
    assert window.project is current
    assert (invalid / "project.json").read_text("utf-8") == "broken"


def test_locked_export_rolls_back_all_previous_files(export_window, qtbot, monkeypatch):
    from app.ui import export_job
    window, _, errors = export_window
    destination = window.project.exports_dir
    names = ["summary.csv", "sites.csv", "annotated_images.zip"]
    for name in names:
        (destination / name).write_bytes(("old " + name).encode())
    replace = export_job.os.replace

    def locked_replace(source, target):
        if target == destination / "sites.csv":
            raise PermissionError("Excel holds sites.csv")
        return replace(source, target)

    monkeypatch.setattr(export_job.os, "replace", locked_replace)
    window.export_csv()
    qtbot.waitUntil(lambda: not window._export_busy, timeout=15000)
    assert errors and "原有导出文件已保留" in errors[-1]
    for name in names:
        assert (destination / name).read_bytes() == ("old " + name).encode()
    assert not list(destination.glob(".export-*"))
