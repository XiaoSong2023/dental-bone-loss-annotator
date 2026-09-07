"""Persistence/integration fixtures; these tests do not demonstrate model accuracy."""
import csv
import copy
import json
from pathlib import Path
import tempfile
import unittest

from app.core import export, storage
from app.core.models import CaseAnnotation, Pt, SCHEMA_VERSION
from app.core.project import CaseEntry, Project


class BoxCoordinatesTests(unittest.TestCase):
    def test_manual_box_four_corners_survive_save_and_follow_reframing(self):
        ann = CaseAnnotation(image_width=1000, image_height=500)
        tooth = ann.add_tooth(16, (0., 12.5, 42.25, 50.5))
        expected = [[0., 12.5], [42.25, 12.5], [42.25, 63.], [0., 63.]]
        self.assertEqual(list(tooth.roi_corners.values()), expected)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "中文标注.json"
            storage.save_annotation(path, ann)
            written = json.loads(path.read_text(encoding="utf-8"))
            self.assertEqual(written["teeth"][0]["roi_corners"], tooth.roi_corners)
            restored = storage.load_annotation(path)
            self.assertEqual(restored.teeth[0].roi_corners, tooth.roi_corners)
        tooth.roi = (2., 3., 10., 20.)
        self.assertEqual(tooth.roi_corners["bottom_right"], [12., 23.])

    def test_legacy_box_derives_corners_without_changing_points(self):
        old = {"schema_version": 3, "teeth": [{"roi": [5, 6, 20, 30],
               "sites": [{"surface": "mesial", "cej": {"x": 0, "y": 0}}]}]}
        migrated = storage.migrate(old)
        self.assertEqual(migrated["schema_version"], SCHEMA_VERSION)
        tooth = CaseAnnotation.from_dict(migrated).teeth[0]
        self.assertEqual(tooth.roi_corners["top_right"], [25., 6.])
        self.assertEqual(tooth.site("mesial").cej.as_tuple(), (0., 0.))

    def test_exports_contain_only_current_landmarks_and_keep_internal_boxes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project = Project.create(root / "项目")
            entry = CaseEntry(case_id="框坐标", image="fixture.png")
            project.cases.append(entry)
            ann = CaseAnnotation(image_width=100, image_height=100)
            marked = ann.add_tooth(16, (1.25, 2.5, 30., 40.))
            marked.site("mesial").set("cej", Pt(8.123456789, 9.5))
            marked.site("mesial").point_meta["cej"]["raw_xy"] = [1., 2.]
            ann.add_tooth(15, (40., 2., 20., 40.))
            project.save_annotation(entry, ann)
            paths = export.export_csv(project, project.grading)
            self.assertEqual([p.name for p in paths], ["summary.csv", "sites.csv"])
            with paths[1].open(encoding="utf-8-sig", newline="") as stream:
                row, = csv.DictReader(stream)
            self.assertEqual(float(row["CEJ_x"]), 8.12)
            self.assertFalse(any("框" in key or "原始" in key for key in row))
            from openpyxl import load_workbook
            book = load_workbook(export.export_excel(project, project.grading))
            self.assertEqual(book.sheetnames, ["患者级汇总", "位点级明细"])
            self.assertEqual(book["位点级明细"].max_row, 2)
            book.close()
            data = json.loads(export.export_single_json(project, entry, root / "导出.json").read_text(encoding="utf-8"))
            self.assertEqual(len(data["teeth"]), 1)
            self.assertEqual(data["teeth"][0]["sites"][0]["cej"]["x"], 8.123456789)
            forbidden = {"roi", "roi_corners", "initial_auto_roi", "segmentation", "contours", "box_history", "ai_runs", "candidate_surfaces", "point_meta", "bone_crest_curve"}
            def keys(value):
                if isinstance(value, dict):
                    for key, child in value.items():
                        yield key
                        yield from keys(child)
                elif isinstance(value, list):
                    for child in value:
                        yield from keys(child)
            self.assertFalse(forbidden.intersection(keys(data)))
            restored = project.load_annotation(entry)
            self.assertEqual(len(restored.teeth), 2)
            self.assertEqual(restored.teeth[0].roi_corners["bottom_right"], [31.25, 42.5])
            self.assertEqual(restored.teeth[0].site("mesial").point_meta["cej"]["raw_xy"], [1., 2.])

    def test_csv_replaces_only_recognised_legacy_box_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            project = Project.create(Path(tmp) / "project")
            legacy = project.exports_dir / "boxes.csv"
            old_header = ["病例号", "片名", "牙序", "牙位FDI",
                          *[f"框{c}_{a}" for c in ("左上", "右上", "右下", "左下") for a in ("x", "y")],
                          "坐标空间", "框来源"]
            export.write_csv(legacy, old_header, [])
            export.export_csv(project, project.grading)
            self.assertFalse(legacy.exists())
            legacy.write_text("user-owned data", encoding="utf-8")
            export.export_csv(project, project.grading)
            self.assertEqual(legacy.read_text(encoding="utf-8"), "user-owned data")


def protocol_fixture():
    """Explicit engineering fixture, never used by application inference."""
    def surface(x):
        return {"object_score": .8, "side": "unknown", "points": {
            "cej": {"xy": [x, 50.], "score": .9},
            "bone_level": {"xy": [x+1, 65.], "score": .8},
            "apex": {"xy": [x+2, 120.], "score": .7}}}
    return {"run_id": "engineering-fixture", "status": "completed", "image": {"width": 200, "height": 200},
            "models": {"pose": {"id": "fixture-not-a-trained-model"}},
            "result_ref": "runs/fixture/output/result.json", "teeth": [
                {"instance_id": "fixture-1", "fdi": 16, "score": .9, "bbox_xyxy": [10., 20., 80., 160.],
                 "contours": [[[10., 20.], [80., 20.], [80., 160.]]], "mask_ref": "masks/one.png",
                 "surfaces": [surface(18.), surface(65.)]},
                {"instance_id": "fixture-2", "fdi": 15, "score": .7, "bbox_xyxy": [85., 25., 120., 140.],
                 "contours": [], "surfaces": []}]}


class AutomaticMergeTests(unittest.TestCase):
    def test_three_model_points_keep_extra_point_missing_and_preserve_raw_after_drag(self):
        from app.core.ai_merge import merge_ai_result
        ann = CaseAnnotation(image_width=200, image_height=200)
        result = protocol_fixture()
        before = copy.deepcopy(result)
        counts = merge_ai_result(ann, result)
        self.assertEqual(counts["points_added"], 6)
        self.assertEqual(counts["teeth_without_points"], 1)
        tooth = ann.teeth[0]
        site = tooth.site("mesial")
        self.assertIsNone(site.cej2)
        self.assertFalse(site.side_confirmed)
        self.assertIn("图像左", site.display_label)
        raw = site.point_meta["cej"]["raw_xy"][:]
        site.set("cej", Pt(20.25123456789, 51.5))
        self.assertEqual(site.point_meta["cej"]["raw_xy"], raw)
        self.assertTrue(site.point_meta["cej"]["edited"])
        restored = CaseAnnotation.from_dict(ann.to_dict())
        self.assertEqual(restored.teeth[0].site("mesial").cej.as_tuple(), (20.25123456789, 51.5))
        self.assertEqual(restored.teeth[0].site("mesial").point_meta["cej"]["raw_xy"], raw)
        self.assertEqual(result, before)
        self.assertEqual(len(restored.ai_runs), 1)

    def test_existing_manual_marks_and_box_survive_rerun_without_duplicate_import(self):
        from app.core.ai_merge import merge_ai_result
        ann = CaseAnnotation(image_width=200, image_height=200)
        tooth = ann.add_tooth(16, (9., 19., 72., 142.))
        tooth.site("mesial").set("cej", Pt(12.5, 44.25))
        result = protocol_fixture()
        merge_ai_result(ann, result)
        self.assertEqual(ann.teeth[0].roi, (9., 19., 72., 142.))
        self.assertEqual(ann.teeth[0].site("mesial").cej.as_tuple(), (12.5, 44.25))
        merge_ai_result(ann, result)
        self.assertEqual(len(ann.teeth), 2)
        self.assertEqual(len(ann.ai_runs), 1)

    def test_invalid_model_coordinate_does_not_partially_modify_case(self):
        from app.core.ai_merge import merge_ai_result
        ann = CaseAnnotation(image_width=200, image_height=200)
        before = ann.to_dict()
        result = protocol_fixture()
        result["teeth"][1]["bbox_xyxy"][0] = float("nan")
        with self.assertRaises(ValueError):
            merge_ai_result(ann, result)
        self.assertEqual(ann.to_dict(), before)

    def test_rerun_keeps_a_manually_rejected_candidate_empty(self):
        from app.core.ai_merge import merge_ai_result
        ann = CaseAnnotation(image_width=200, image_height=200)
        result = protocol_fixture()
        merge_ai_result(ann, result)
        site = ann.teeth[0].site("mesial")
        site.set("cej", None)
        rejected_metadata = copy.deepcopy(site.point_meta["cej"])
        result["run_id"] = "second-engineering-run"
        merge_ai_result(ann, result)
        restored = ann.teeth[0].site("mesial")
        self.assertIsNone(restored.cej)
        self.assertEqual(restored.point_meta["cej"], rejected_metadata)

    def test_larger_manual_context_box_matches_same_fdi_without_losing_box(self):
        from app.core.ai_merge import merge_ai_result
        ann = CaseAnnotation(image_width=200, image_height=200)
        ann.add_tooth(16, (0., 0., 100., 200.))
        result = protocol_fixture()
        result["teeth"] = result["teeth"][:1]
        result["teeth"][0]["bbox_xyxy"] = [20., 40., 80., 160.]
        counts = merge_ai_result(ann, result)
        self.assertEqual(counts["matched_teeth"], 1)
        self.assertEqual(len(ann.teeth), 1)
        self.assertEqual(ann.teeth[0].roi, (0., 0., 100., 200.))

        different_fdi = CaseAnnotation(image_width=200, image_height=200)
        different_fdi.add_tooth(17, (0., 0., 100., 200.))
        merge_ai_result(different_fdi, result)
        self.assertEqual(len(different_fdi.teeth), 2)


if __name__ == "__main__":
    unittest.main()
