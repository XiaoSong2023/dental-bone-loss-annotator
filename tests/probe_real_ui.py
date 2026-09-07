"""Explicit, opt-in UI verification with an existing REAL inference result.

This script never loads a model and is not discovered by pytest. It creates a
separate review project; the mouse edit is an engineering test, not a clinical
correction. Example: python tests/probe_real_ui.py --result .../result.json
--image .../panoramic.jpg --output reports/real_ui_probe.local
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--image", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result_path, image_path, output = args.result.resolve(), args.image.resolve(), args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    report = {
        "started_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "running", "new_inference_calls": 0,
        "inference_protocol_exercised": False,
        "scope": "Existing real inference result replay through UI completion, editing and exports; not a complete QProcess run",
        "source_result": str(result_path), "source_image": str(image_path),
        "mouse_edit_is_engineering_verification_not_clinical_review": True,
        "assertions": [],
    }

    def check(name, condition, **details):
        report["assertions"].append({"name": name, "passed": bool(condition), **details})
        if not condition:
            raise AssertionError(name + ": " + str(details))

    from PySide6 import __version__ as qt_version
    from PySide6.QtCore import QPoint, QPointF, QRectF, QSettings, Qt
    from PySide6.QtGui import QFont, QFontDatabase
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
    from openpyxl import load_workbook
    from app.core.project import Project
    from app.ui import main_window, theme
    from app.ui.widgets.point_item import PointItem
    from app.ui.tooth_dialog import ToothDialog
    from app.core.formula import get
    from app.core.review_geometry import project_to_contours, edit_box, undo_box

    app = QApplication.instance() or QApplication([])
    app.setQuitOnLastWindowClosed(False)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if font_path.is_file():
        font_id = QFontDatabase.addApplicationFont(str(font_path))
        families = QFontDatabase.applicationFontFamilies(font_id)
        if families:
            app.setFont(QFont(families[0], 9))
    theme.apply(app)
    report["environment"] = {"python": sys.version, "pyside6": qt_version,
                             "qt_platform": app.platformName()}
    window = None
    notices = []

    def fail_dialog(_parent, title, text, *_args, **_kwargs):
        raise RuntimeError(f"{title}: {text}")

    def settle():
        for _ in range(5):
            app.processEvents()

    try:
        original = json.loads(result_path.read_text(encoding="utf-8"))
        check("source_hash_matches_real_model_input", digest(image_path) == original["image"]["sha256"])
        project = Project.create(output / "review_project")
        imported, _ = project.import_paths([image_path])
        check("one_real_panorama_imported", len(imported) == 1)
        run_dir = project.root / "runs" / original["run_id"]
        copied = run_dir / "output"
        shutil.copytree(result_path.parent, copied)
        copied_result = copied / result_path.name
        input_snapshot = run_dir / ("source-input" + image_path.suffix)
        shutil.copy2(image_path, input_snapshot)
        check("original_model_output_copied_unchanged", digest(copied_result) == digest(result_path))
        check("input_snapshot_matches_unmodified_model_source", digest(input_snapshot) == original["image"]["sha256"])

        with patch.object(main_window.MainWindow, "_restore_geometry", lambda self: None), \
             patch.object(main_window, "QSettings", lambda *_args:
                          QSettings(str(output / "probe-settings.ini"), QSettings.Format.IniFormat)), \
             patch.object(QMessageBox, "information", lambda _p, title, text, *_a: notices.append([title, text])), \
             patch.object(QMessageBox, "warning", fail_dialog), \
             patch.object(QMessageBox, "critical", fail_dialog):
            window = main_window.MainWindow()
            window.resize(1560, 940)
            window.show()
            window._load_project(project)
            settle()
            check("image_coordinate_dimensions_match", (window.ann.image_width, window.ann.image_height) ==
                  (original["image"]["width"], original["image"]["height"]))
            window._ai_context = (window.project, window.entry, window.ann, run_dir,
                                  input_snapshot, "source_image_snapshot_no_conversion")
            window._ai_busy = True
            window.btn_pick.setChecked(False)
            window._update_enabled()
            window._ai_completed(str(copied_result))
            settle()
            ann = window.ann
            selected_points = [(t, s, role, s.get(role)) for t in ann.teeth for s in t.sites
                               for role in ("cej", "ac", "apex") if s.get(role) is not None]
            raw_surfaces = sum(len(t.get("surfaces", [])) for t in original["teeth"])
            report["model_and_ui_counts"] = {
                "model_teeth": len(original["teeth"]), "raw_model_surfaces": raw_surfaces,
                "raw_model_points": sum(sum(p is not None for p in s.get("points", {}).values())
                                        for t in original["teeth"] for s in t.get("surfaces", [])),
                "raw_teeth_with_surfaces": sum(bool(t.get("surfaces")) for t in original["teeth"]),
                "ui_teeth": len(ann.teeth), "ui_surfaces": sum(bool(s.filled_count) for t in ann.teeth for s in t.sites),
                "ui_points": len(selected_points),
                "ui_teeth_with_points": sum(any(s.filled_count for s in t.sites) for t in ann.teeth),
                "merge_counts": ann.ai_runs[-1]["counts"],
            }
            check("all_real_teeth_applied", len(ann.teeth) == len(original["teeth"]))
            check("three_point_candidates_present", len(selected_points) > 0)
            check("cej2_not_invented", all(s.cej2 is None for t in ann.teeth for s in t.sites))
            check("anatomical_sides_not_assumed", all(not s.side_confirmed for t in ann.teeth for s in t.sites if s.filled_count))
            check("run_reference_is_portable", not Path(ann.ai_runs[-1]["result_ref"]).is_absolute()
                  and (project.root / ann.ai_runs[-1]["result_ref"]).is_file())
            check("analysis_input_provenance_is_accurate", ann.ai_runs[-1]["analysis_input"]["sha256"] == original["image"]["sha256"]
                  and ann.ai_runs[-1]["analysis_input"]["preconversion"] == "source_image_snapshot_no_conversion"
                  and ann.ai_runs[-1]["source_case"]["sha1"] == window.entry.sha1)
            check("all_segmentation_masks_preserved", all((project.root / t.segmentation["mask_ref"]).is_file() for t in ann.teeth))
            residuals = [project_to_contours(p.as_tuple(), t.segmentation["contours"])[1]["distance_px"]
                         for t, s, role, p in selected_points]
            check("all_displayed_candidates_on_displayed_contour", max(residuals) < 1e-6,
                  max_residual_px=max(residuals))
            original_raw = [tuple(p.get("raw_xy", p["xy"])) for t in original["teeth"]
                            for s in t["surfaces"] for p in s["points"].values() if p]
            check("all_displayed_candidates_keep_actual_model_raw", all(
                tuple(s.point_meta[role]["raw_xy"]) in original_raw for t, s, role, p in selected_points))
            window.setWindowTitle("真实模型自动结果 · UI 验证测试项目")
            window.view.fit()
            settle()
            window.grab().save(str(output / "01-automatic-result.png"))

            # Pick the most isolated real apex so the Qt hit test is unambiguous.
            choices = [entry for entry in selected_points if entry[2] == "apex"]
            tooth, site, role, point = max(choices, key=lambda candidate:
                min(math.hypot(candidate[3].x - other[3].x, candidate[3].y - other[3].y)
                    for other in selected_points if other is not candidate))
            tooth_id, surface = tooth.tooth_id, site.surface
            before = point.as_tuple()
            raw = site.point_meta[role]["raw_xy"][:]
            window.view.fit(QRectF(point.x - 90, point.y - 90, 180, 180))
            settle()
            start = window.view.mapFromScene(QPointF(*before))
            end = start + QPoint(24, 18)
            clicked_item = window.view.itemAt(start)
            check("mouse_hit_is_real_candidate_point", isinstance(clicked_item, PointItem),
                  hit_type=type(clicked_item).__name__)
            start_scene, end_scene = window.view.mapToScene(start), window.view.mapToScene(end)
            expected = (before[0] + end_scene.x() - start_scene.x(), before[1] + end_scene.y() - start_scene.y())
            viewport = window.view.viewport()
            QTest.mousePress(viewport, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, start)
            QTest.mouseMove(viewport, (start + end) / 2)
            settle()
            QTest.mouseMove(viewport, end)
            settle()
            QTest.mouseRelease(viewport, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, end)
            settle()
            after = site.get(role).as_tuple()
            check("real_qt_mouse_drag_updates_original_coordinates", math.dist(before, after) > 1 and math.dist(expected, after) < .1,
                  before=before, after=after, expected=expected)
            check("drag_keeps_raw_model_point", site.point_meta[role]["raw_xy"] == raw)
            check("drag_marks_review_without_confirming_side", site.point_meta[role]["edited"] and
                  site.point_meta[role]["review_status"] == "reviewed" and not site.side_confirmed)
            report["mouse_drag"] = {"tooth_id": tooth_id, "fdi": tooth.fdi, "slot": surface, "role": role,
                                    "raw_xy": raw, "before_xy": before, "after_xy": after,
                                    "viewport_displacement": [24, 18]}
            window.grab().save(str(output / "02-mouse-drag.png"))
            window.save_current(silent=True)
            window._load_project(Project.open(project.root))
            settle()
            restored_tooth = next(t for t in window.ann.teeth if t.tooth_id == tooth_id)
            restored_site = restored_tooth.site(surface)
            check("reopen_retains_manual_edit", math.dist(restored_site.get(role).as_tuple(), after) < 1e-5)
            check("reopen_retains_raw_prediction", restored_site.point_meta[role]["raw_xy"] == raw)
            check("reopen_retains_all_teeth", len(window.ann.teeth) == len(original["teeth"]))
            original_box = restored_tooth.roi
            x, y, w, h = original_box
            edit_box(window.ann, (x+1, y+1, w-2, h-2), restored_tooth)
            window.save_current(silent=True)
            window._load_project(Project.open(project.root))
            check("box_undo_history_survives_reopen", bool(window.ann.box_history))
            check("box_undo_succeeds_after_reopen", undo_box(window.ann))
            restored_tooth = next(t for t in window.ann.teeth if t.tooth_id == tooth_id)
            restored_site = restored_tooth.site(surface)
            check("box_undo_recovers_real_auto_box_and_keeps_drag", restored_tooth.roi == original_box
                  and math.dist(restored_site.get(role).as_tuple(), after) < 1e-5)
            window.save_current(silent=True)
            window.btn_pick.setChecked(False)
            window.view.fit()
            settle()
            window.grab().save(str(output / "03-reopened-project.png"))

            dialog = ToothDialog(window._gray, window.ann, [restored_tooth],
                                 get(window.ann.formula_id), project.grading.formula_config(), window)
            dialog.show()
            settle()
            dialog.grab().save(str(output / "04-review-dialog.png"))
            dialog.reject()

            window.export_csv()
            window.export_excel()
            exported_json = project.exports_dir / "real-case.reviewed.json"
            with patch.object(QFileDialog, "getSaveFileName", return_value=(str(exported_json), "JSON (*.json)")):
                window.export_json()
            data = json.loads(exported_json.read_text(encoding="utf-8"))
            exported_tooth = next(t for t in data["teeth"] if t["tooth_id"] == tooth_id)
            exported_site = next(s for s in exported_tooth["sites"] if s["surface"] == surface)
            check("json_retains_edited_and_raw_coordinates", math.dist([exported_site[role]["x"], exported_site[role]["y"]], after) < 1e-5
                  and exported_site["point_meta"][role]["raw_xy"] == raw)
            with (project.exports_dir / "sites.csv").open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            row = next(row for row in rows if int(row["牙序"]) == restored_tooth.index+1 and row["邻面"] == restored_site.display_label)
            check("csv_exports_manual_edit_and_original_prediction", math.dist([float(row["APEX_x"]), float(row["APEX_y"])], after) < .008
                  and math.dist([float(row["APEX_原始x"]), float(row["APEX_原始y"])], raw) < 1e-5
                  and row["APEX_人工修正"] == "是" and row["侧别已确认"] == "否")
            with (project.exports_dir / "boxes.csv").open(encoding="utf-8-sig", newline="") as stream:
                boxes = list(csv.DictReader(stream))
            check("csv_exports_all_four_box_corners", len(boxes) == len(original["teeth"]) and
                  all(all(row[f"框{corner}_{axis}"] != "" for corner in ("左上", "右上", "右下", "左下")
                          for axis in ("x", "y")) for row in boxes))
            workbook = load_workbook(project.exports_dir / "汇总.xlsx", read_only=True, data_only=True)
            check("excel_exports_sites_and_boxes", workbook["位点级明细"].max_row == len(rows)+1 and
                  workbook["牙框坐标"].max_row == len(boxes)+1)
            values = list(workbook["位点级明细"].values)
            excel_rows = [dict(zip(values[0], row)) for row in values[1:]]
            excel_row = next(row for row in excel_rows if int(row["牙序"]) == restored_tooth.index+1 and row["邻面"] == restored_site.display_label)
            check("excel_retains_adjusted_apex", math.dist([float(excel_row["APEX_x"]), float(excel_row["APEX_y"])], after) < .008)
            workbook.close()
            report["exported_files"] = [{"path": str(path.relative_to(output)), "sha256": digest(path), "bytes": path.stat().st_size}
                                        for path in sorted(project.exports_dir.iterdir()) if path.is_file()]
            report["ui_notices"] = notices
            report["status"] = "passed"
            window.close()
            window = None
    except Exception as exc:
        report["status"] = "failed"
        report["error"] = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        report["finished_at_utc"] = datetime.now(timezone.utc).isoformat()
        (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        if window is not None:
            window._ai.stop()
        print(json.dumps({"status": report["status"], "report": str(output / "report.json")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
