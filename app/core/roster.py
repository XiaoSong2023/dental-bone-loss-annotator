"""名册导入 —— 从 CSV / Excel 批量带入年龄、性别、备注。

年龄是算「骨吸收年龄比」和 Grade 的必需数据，70 多个病例逐个手填太慢。
列名中英文、大小写、常见别名都能自动识别。
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

# 列名别名表（全部转小写、去空格后比对）
ALIASES: dict[str, tuple[str, ...]] = {
    "key": ("文件名", "档案名", "图片", "图片名", "片子", "filename", "file", "file_name", "image", "img"),
    "id": ("病例号", "编号", "序号", "患者编号", "id", "case_id", "patient_id", "no", "number"),
    "age": ("年龄", "岁", "age", "years"),
    "sex": ("性别", "sex", "gender"),
    "note": ("备注", "说明", "note", "notes", "remark", "comment"),
}


@dataclass
class RosterRow:
    key: str = ""
    id: str = ""
    age: float | None = None
    sex: str = ""
    note: str = ""

    @property
    def match_keys(self) -> list[str]:
        """可用于匹配的候选键：文件名（含/不含扩展名）与病例号。"""
        out = []
        for raw in (self.key, self.id):
            v = (raw or "").strip()
            if not v:
                continue
            out.append(v)
            stem = Path(v).stem
            if stem != v:
                out.append(stem)
        return out


@dataclass
class MatchReport:
    matched: list[tuple[str, RosterRow]] = field(default_factory=list)   # (case_id, row)
    unmatched_cases: list[str] = field(default_factory=list)             # 项目里没对上名册的
    unused_rows: list[RosterRow] = field(default_factory=list)           # 名册里多出来的
    errors: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return (
            f"成功匹配 {len(self.matched)} 例 · "
            f"未匹配 {len(self.unmatched_cases)} 例 · "
            f"名册多余 {len(self.unused_rows)} 行"
        )


def _canon(s: str) -> str:
    return str(s or "").strip().lower().replace(" ", "").replace("　", "")


def _map_headers(headers: Iterable[str]) -> dict[int, str]:
    """表头 → 字段名。返回 {列序号: 字段}。"""
    out: dict[int, str] = {}
    for i, h in enumerate(headers):
        c = _canon(h)
        if not c:
            continue
        for field_name, names in ALIASES.items():
            if c in names or any(c == _canon(n) for n in names):
                out.setdefault(i, field_name)
                break
        else:
            # 宽松匹配：表头里“含”别名即可，如「患者年龄(岁)」
            for field_name, names in ALIASES.items():
                if any(_canon(n) and _canon(n) in c for n in names):
                    out.setdefault(i, field_name)
                    break
    return out


def _parse_age(v) -> float | None:
    if v in (None, ""):
        return None
    s = str(v).strip()
    # 容忍「45岁」「45 y」这类写法
    num = "".join(ch for ch in s if ch.isdigit() or ch == ".")
    try:
        age = float(num)
    except ValueError:
        return None
    return age if 0 < age < 130 else None


def _rows_from_csv(path: Path) -> list[list[str]]:
    # 中文 Excel 导出的 CSV 常是 GBK，且可能带 BOM
    for enc in ("utf-8-sig", "gbk", "utf-8", "latin-1"):
        try:
            with path.open("r", encoding=enc, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                return [row for row in csv.reader(f, dialect)]
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError("无法识别 CSV 的文字编码")


def _rows_from_excel(path: Path) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError("读取 Excel 需要 openpyxl，请先安装") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def read_roster(path: Path) -> list[RosterRow]:
    path = Path(path)
    raw = _rows_from_excel(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _rows_from_csv(path)
    raw = [r for r in raw if any(str(c).strip() for c in r if c is not None)]
    if not raw:
        return []

    mapping = _map_headers(str(c) if c is not None else "" for c in raw[0])
    if not mapping:
        raise ValueError(
            "认不出表头。至少需要一列文件名或病例号，外加一列年龄。\n"
            "可识别的列名：文件名/病例号/编号/年龄/性别/备注（中英文皆可）"
        )

    out: list[RosterRow] = []
    for r in raw[1:]:
        row = RosterRow()
        for i, fname in mapping.items():
            if i >= len(r):
                continue
            v = r[i]
            if v is None:
                continue
            if fname == "age":
                row.age = _parse_age(v)
            else:
                setattr(row, fname, str(v).strip())
        if row.match_keys:
            out.append(row)
    return out


def match(case_ids: Iterable[str], rows: Iterable[RosterRow]) -> MatchReport:
    """名册行 ↔ 项目病例。匹配键大小写不敏感，文件名可带可不带扩展名。"""
    rows = list(rows)
    case_ids = list(case_ids)

    index: dict[str, RosterRow] = {}
    for row in rows:
        for k in row.match_keys:
            index.setdefault(_canon(k), row)

    report = MatchReport()
    used: set[int] = set()
    for cid in case_ids:
        row = index.get(_canon(cid)) or index.get(_canon(Path(cid).stem))
        if row is None:
            report.unmatched_cases.append(cid)
        else:
            report.matched.append((cid, row))
            used.add(id(row))
    report.unused_rows = [r for r in rows if id(r) not in used]
    return report
